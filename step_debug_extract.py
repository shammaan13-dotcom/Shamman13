from openpyxl import load_workbook, Workbook
import os

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------

RAMIS_FILE = os.getenv("RAMIS_FILE", "input/ramis.xlsx")
MACL_FILE  = os.getenv("MACL_FILE", "input/macl_master.xlsx")

OUTPUT_FILE = "output/DEBUG_EXTRACTION.xlsx"

os.makedirs("output", exist_ok=True)


# ------------------------------------------------------------
# LOAD FILES
# ------------------------------------------------------------

ramis_wb = load_workbook(RAMIS_FILE, data_only=True)
macl_wb  = load_workbook(MACL_FILE)

ramis_ws = ramis_wb.active
macl_ws  = macl_wb.active

print("FILES LOADED")


# ------------------------------------------------------------
# CREATE OUTPUT
# ------------------------------------------------------------

out_wb = Workbook()

macl_out = out_wb.active
macl_out.title = "MACL_DATA"

ramis_out = out_wb.create_sheet("RAMIS_DATA")


# ------------------------------------------------------------
# COPY MACL (A–G)
# ------------------------------------------------------------

print("COPYING MACL...")

row_out = 1

for r in range(1, macl_ws.max_row + 1):

    row_vals = []

    for c in range(1, 8):  # A–G
        row_vals.append(macl_ws.cell(r, c).value)

    if any(row_vals):
        for c, val in enumerate(row_vals, start=1):
            macl_out.cell(row_out, c).value = val

        row_out += 1


# ------------------------------------------------------------
# COPY RAMIS (A–F)
# ------------------------------------------------------------

print("COPYING RAMIS...")

row_out = 1

for r in range(1, ramis_ws.max_row + 1):

    row_vals = []

    for c in range(1, 7):  # A–F
        row_vals.append(ramis_ws.cell(r, c).value)

    if any(row_vals):
        for c, val in enumerate(row_vals, start=1):
            ramis_out.cell(row_out, c).value = val

        row_out += 1


# ------------------------------------------------------------
# SAVE
# ------------------------------------------------------------

out_wb.save(OUTPUT_FILE)

print("DEBUG FILE CREATED:", OUTPUT_FILE)
