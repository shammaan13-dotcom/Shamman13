# =================================================
# FINAL RECONCILIATION ENGINE (LOCAL VERSION)
# =================================================

import pandas as pd
from datetime import datetime, timedelta, time
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os


# ------------------------------------------------
# CONFIG
# ------------------------------------------------

INPUT_FILE = "output/MASTER_MACL_UPDATED.xlsx"
OUTPUT_FILE = "output/FINAL_OUTPUT.xlsx"


# ------------------------------------------------
# CHECK FILE
# ------------------------------------------------

if not os.path.exists(INPUT_FILE):
    raise FileNotFoundError(f"Missing file: {INPUT_FILE}")


# ------------------------------------------------
# LOAD
# ------------------------------------------------

df = pd.read_excel(INPUT_FILE, header=None)
out = df.copy()


# ------------------------------------------------
# COLUMN INDEX
# ------------------------------------------------

MACL_AIRLINE = 0
MACL_DAY = 1
MACL_FLT = 3
MACL_STA = 4
MACL_STD = 5
MACL_EFF = 6

RAMIS_AIRLINE = 8
RAMIS_DAY = 9
RAMIS_FLT = 10
RAMIS_STA = 11
RAMIS_STD = 12
RAMIS_EFF = 13

RAMIS_START = 8

NEW_FLT = 14
NEW_STA = 15
NEW_STD = 16
NEW_EFF = 17
STATUS  = 18
COMMENT = 19


# ------------------------------------------------
# TIME HELPERS
# ------------------------------------------------

def to_time(v):
    try:
        return datetime.strptime(str(v), "%H:%M").time()
    except:
        return None

def base_flt(f):
    return str(f).split("-")[0]

STA_START = time(4,45)
STA_END   = time(15,45)

STD_START = time(9,0)
STD_END   = time(23,59)


# ------------------------------------------------
# BUILD RAMIS LOOKUP
# ------------------------------------------------

ramis_rows = {}

for i,r in df.iterrows():

    airline = str(r[RAMIS_AIRLINE]).strip().upper()
    day     = str(r[RAMIS_DAY]).strip().upper()
    flt     = str(r[RAMIS_FLT]).strip().upper()
    sta     = to_time(r[RAMIS_STA])
    std     = to_time(r[RAMIS_STD])

    if airline and day and flt and flt != "NAN":
        ramis_rows.setdefault(airline, []).append((day,flt,sta,std,i))

used_ramis = set()


# ------------------------------------------------
# MATCHING
# ------------------------------------------------

for i,r in df.iterrows():

    airline  = str(r[MACL_AIRLINE]).strip().upper()
    macl_day = str(r[MACL_DAY]).strip().upper()
    macl_flt = str(r[MACL_FLT]).strip().upper()

    macl_sta = to_time(r[MACL_STA])
    macl_std = to_time(r[MACL_STD])

    macl_norm = macl_flt.replace("-","")

    found = False

    for r_day,r_flt,r_sta,r_std,r_i in ramis_rows.get(airline,[]):

        if r_i in used_ramis:
            continue

        if macl_day != r_day:
            continue

        if macl_norm == r_flt.replace("-",""):
            out.iloc[i,RAMIS_START:] = df.iloc[r_i,RAMIS_START:]
            used_ramis.add(r_i)
            found = True
            break

    if not found:
        out.iloc[i,RAMIS_START:] = ""


# ------------------------------------------------
# RECONCILIATION LOGIC
# ------------------------------------------------

for i,r in out.iterrows():

    macl_air = str(r[MACL_AIRLINE]).upper()

    if "CARGO" in macl_air:
        out.iloc[i,STATUS] = "IGNORED"
        out.iloc[i,COMMENT] = "CARGO"
        continue

    macl_eff = str(r[MACL_EFF])
    ramis_eff = str(r[RAMIS_EFF])

    if macl_eff != ramis_eff:
        out.iloc[i,NEW_EFF] = macl_eff
        out.iloc[i,STATUS] = "CHANGE"
        out.iloc[i,COMMENT] = "EFFECTIVE CHANGE"
    else:
        out.iloc[i,STATUS] = "NO CHANGE"
        out.iloc[i,COMMENT] = "-"


# ------------------------------------------------
# EXPORT
# ------------------------------------------------

out.to_excel(OUTPUT_FILE,index=False,header=False)

wb = load_workbook(OUTPUT_FILE)
ws = wb.active


# ------------------------------------------------
# BASIC FORMATTING
# ------------------------------------------------

ws.freeze_panes = "A3"

for col in ws.columns:

    max_len = 0
    col_letter = get_column_letter(col[0].column)

    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))

    ws.column_dimensions[col_letter].width = max_len + 2


# ------------------------------------------------
# SAVE OUTPUT
# ------------------------------------------------

import os

OUTPUT_FILE = "output/FINAL_OUTPUT.xlsx"

# Create output folder if not exists
os.makedirs("output", exist_ok=True)

# Save workbook
wb.save(OUTPUT_FILE)

print(f"✅ FINAL OUTPUT READY: {OUTPUT_FILE}")
