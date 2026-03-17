import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from collections import defaultdict
from datetime import datetime
from io import BytesIO


def process(macl_file, ramis_file):

    # -----------------------------
    # LOAD FILES FROM STREAMLIT
    # -----------------------------
    ramis_wb = load_workbook(ramis_file, data_only=True)
    master_wb = load_workbook(macl_file)

    ramis_ws = ramis_wb.active
    target_ws = master_wb.active

    # -----------------------------
    # TIME WINDOWS
    # -----------------------------
    ARR_START = datetime.strptime("04:45","%H:%M").time()
    ARR_END   = datetime.strptime("15:45","%H:%M").time()

    DEP_START = datetime.strptime("09:00","%H:%M").time()
    DEP_END   = datetime.strptime("23:59","%H:%M").time()

    # -----------------------------
    # HELPERS
    # -----------------------------
    def parse_time(t):
        try:
            return datetime.strptime(str(t), "%H:%M").time()
        except:
            return None

    def correct_midnight_std(std):
        if std and (std.hour == 0 or std.hour == 1):
            return datetime.strptime("23:59","%H:%M").time()
        return std

    def validate_flight(flt, sta, std):

        sta_t = parse_time(sta)
        std_t = correct_midnight_std(parse_time(std))

        arr_valid = sta_t and ARR_START <= sta_t <= ARR_END
        dep_valid = std_t and DEP_START <= std_t <= DEP_END

        if not arr_valid and not dep_valid:
            return None, None, None

        if arr_valid and dep_valid:
            if "-" not in flt:
                flt_out = flt.replace("D","") + "-D"
            else:
                flt_out = flt
            return flt_out, sta, std_t.strftime("%H:%M")

        if arr_valid:
            return flt.replace("D",""), sta, ""

        if dep_valid:
            return flt, "", std_t.strftime("%H:%M")

    def normalize_day(day):
        mapping = {
            "MON":"MONDAY","TUE":"TUESDAY","WED":"WEDNESDAY",
            "THU":"THURSDAY","FRI":"FRIDAY","SAT":"SATURDAY","SUN":"SUNDAY"
        }
        return mapping.get(str(day).strip().upper(), None)

    # -----------------------------
    # GROUP DATA
    # -----------------------------
    week_groups = defaultdict(list)

    for row in ramis_ws.iter_rows(min_row=2, max_col=6, values_only=True):

        if not any(row):
            continue

        weekday = normalize_day(row[1])
        if not weekday:
            continue

        airline, day, flt, sta, std, eff = row

        new_flt, new_sta, new_std = validate_flight(flt, sta, std)

        if new_flt is None:
            continue

        week_groups[weekday].append(
            (airline, day, new_flt, new_sta, new_std, eff)
        )

    # -----------------------------
    # SORT
    # -----------------------------
    for d in week_groups:
        week_groups[d] = sorted(week_groups[d], key=lambda x: str(x[0]))

    # -----------------------------
    # STYLES
    # -----------------------------
    purple_fill = PatternFill(start_color="D9CCE3", end_color="D9CCE3", fill_type="solid")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # -----------------------------
    # CLEAR OLD DATA
    # -----------------------------
    for row in range(3, target_ws.max_row + 1):
        for col in range(9, 15):
            target_ws.cell(row=row, column=col).value = None

    # -----------------------------
    # WRITE DATA
    # -----------------------------
    WEEKDAYS = ["MONDAY","TUESDAY","WEDNESDAY","THURSDAY","FRIDAY","SATURDAY","SUNDAY"]

    current_row = 3

    for day in WEEKDAYS:

        if day not in week_groups:
            continue

        target_ws.cell(row=current_row, column=9).value = day

        for col in range(9,15):
            c = target_ws.cell(row=current_row, column=col)
            c.fill = purple_fill
            c.font = header_font
            c.alignment = center

        current_row += 1

        headers = ["AIRLINE","DAYS OF OPS","FLT NO","STA","STD","EFFECTIVE"]

        for i,h in enumerate(headers):
            c = target_ws.cell(row=current_row, column=9+i)
            c.value = h
            c.fill = purple_fill
            c.font = header_font
            c.alignment = center

        current_row += 1

        for record in week_groups[day]:
            for i,val in enumerate(record):
                target_ws.cell(row=current_row, column=9+i).value = val
            current_row += 1

        current_row += 1

    # -----------------------------
    # RETURN FILE
    # -----------------------------
    output = BytesIO()
    master_wb.save(output)
    output.seek(0)

    return output
