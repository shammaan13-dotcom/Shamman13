# ============================================================
# RAMIS PURPLE SIDE – PRODUCTION ENGINE
# WITH TMA STA/STD VALIDATION + MIDNIGHT FIX
# ============================================================

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from collections import defaultdict
from datetime import datetime
from google.colab import files


# ------------------------------------------------------------
# TMA TIME WINDOWS
# ------------------------------------------------------------

ARR_START = datetime.strptime("04:45","%H:%M").time()
ARR_END   = datetime.strptime("15:45","%H:%M").time()

DEP_START = datetime.strptime("09:00","%H:%M").time()
DEP_END   = datetime.strptime("23:59","%H:%M").time()


# ------------------------------------------------------------
# TIME PARSER
# ------------------------------------------------------------

def parse_time(t):

    if t is None:
        return None

    t = str(t).strip()

    if t == "" or t == "-":
        return None

    try:
        return datetime.strptime(t,"%H:%M").time()
    except:
        return None


# ------------------------------------------------------------
# MIDNIGHT STD CORRECTION
# ------------------------------------------------------------

def correct_midnight_std(std):

    if std is None:
        return None

    if std.hour == 0 or std.hour == 1:
        return datetime.strptime("23:59","%H:%M").time()

    return std


# ------------------------------------------------------------
# VALIDATION ENGINE
# ------------------------------------------------------------

def validate_flight(flt, sta, std):

    sta_t = parse_time(sta)
    std_t = parse_time(std)

    std_t = correct_midnight_std(std_t)

    arr_valid = False
    dep_valid = False

    if sta_t:
        if ARR_START <= sta_t <= ARR_END:
            arr_valid = True

    if std_t:
        if DEP_START <= std_t <= DEP_END:
            dep_valid = True

    if not arr_valid and not dep_valid:
        return None,None,None

    # ---------------------------
    # ARR + DEP
    # ---------------------------

    if arr_valid and dep_valid:

        if "-" not in flt:

            if flt.endswith("D"):
                base = flt[:-1]
                flt_out = base + "-D"

            else:
                flt_out = flt + "-D"

        else:
            flt_out = flt

        return flt_out, sta, std_t.strftime("%H:%M")

    # ---------------------------
    # ARR ONLY
    # ---------------------------

    if arr_valid and not dep_valid:

        if "-" in flt:
            flt_out = flt.split("-")[0]
        else:
            flt_out = flt.replace("D","")

        return flt_out, sta, ""

    # ---------------------------
    # DEP ONLY
    # ---------------------------

    if dep_valid and not arr_valid:

        if "-" in flt:
            dep = flt.split("-")[1]
            flt_out = dep
        else:
            flt_out = flt

        return flt_out, "", std_t.strftime("%H:%M")


# ------------------------------------------------------------
# FILE UPLOAD
# ------------------------------------------------------------

print("Upload RAMIS Clean Sheet")
ramis_upload = files.upload()
ramis_file = list(ramis_upload.keys())[0]

print("Upload MASTER MACL WINTER vs RAMIS file")
master_upload = files.upload()
master_file = list(master_upload.keys())[0]

OUTPUT_FILE = "MASTER_MACL_WINTER_vs_RAMIS_UPDATED.xlsx"


# ------------------------------------------------------------
# LOAD FILES
# ------------------------------------------------------------

ramis_wb = load_workbook(ramis_file,data_only=True)
master_wb = load_workbook(master_file)

ramis_ws = ramis_wb.active
target_ws = master_wb.active


# ------------------------------------------------------------
# STYLES
# ------------------------------------------------------------

purple_fill = PatternFill(start_color="D9CCE3",end_color="D9CCE3",fill_type="solid")
header_font = Font(bold=True)
center_align = Alignment(horizontal="center",vertical="center")


# ------------------------------------------------------------
# NORMALIZE DAY
# ------------------------------------------------------------

def normalize_day(day):

    day = str(day).strip().upper()

    mapping = {

        "MON":"MONDAY","MONDAY":"MONDAY",
        "TUE":"TUESDAY","TUESDAY":"TUESDAY",
        "WED":"WEDNESDAY","WEDNESDAY":"WEDNESDAY",
        "THU":"THURSDAY","THURSDAY":"THURSDAY",
        "FRI":"FRIDAY","FRIDAY":"FRIDAY",
        "SAT":"SATURDAY","SATURDAY":"SATURDAY",
        "SUN":"SUNDAY","SUNDAY":"SUNDAY"
    }

    return mapping.get(day,None)


# ------------------------------------------------------------
# GROUP RAMIS DATA BY WEEKDAY
# ------------------------------------------------------------

week_groups = defaultdict(list)

for row in ramis_ws.iter_rows(min_row=2,max_col=6,values_only=True):

    if not any(row):
        continue

    weekday = normalize_day(row[1])

    if not weekday:
        continue

    airline,day,flt,sta,std,eff = row

    validated = validate_flight(flt,sta,std)

    if validated[0] is None:
        continue

    new_flt,new_sta,new_std = validated

    week_groups[weekday].append(

        (
            airline,
            day,
            new_flt,
            new_sta,
            new_std,
            eff
        )
    )


# ------------------------------------------------------------
# SORT BY AIRLINE
# ------------------------------------------------------------

for day in week_groups:

    week_groups[day] = sorted(

        week_groups[day],
        key=lambda x: str(x[0]).strip().upper()
    )


# ------------------------------------------------------------
# CLEAR OLD PURPLE SECTION
# ------------------------------------------------------------

for row in range(3,target_ws.max_row+1):

    for col in range(9,15):

        cell = target_ws.cell(row=row,column=col)

        is_merged = False

        for merged in target_ws.merged_cells.ranges:

            if cell.coordinate in merged:
                is_merged = True
                break

        if not is_merged:
            cell.value = None


# ------------------------------------------------------------
# WEEKDAY ORDER
# ------------------------------------------------------------

WEEKDAYS = [

    "MONDAY","TUESDAY","WEDNESDAY",
    "THURSDAY","FRIDAY","SATURDAY","SUNDAY"
]


# ------------------------------------------------------------
# WRITE PURPLE SECTION
# ------------------------------------------------------------

current_row = 3

for index,day in enumerate(WEEKDAYS):

    if day not in week_groups:
        continue

    if index == 0:

        for record in week_groups[day]:

            for offset,value in enumerate(record):

                target_ws.cell(row=current_row,column=9+offset).value = value

            current_row += 1

    else:

        current_row += 1

        target_ws.cell(row=current_row,column=9).value = day

        for col in range(9,15):

            cell = target_ws.cell(row=current_row,column=col)

            cell.fill = purple_fill
            cell.font = header_font
            cell.alignment = center_align

        current_row += 1

        headers = ["AIRLINE","DAYS OF OPS","FLT NO","STA","STD","EFFECTIVE"]

        for i,header in enumerate(headers):

            cell = target_ws.cell(row=current_row,column=9+i)

            cell.value = header
            cell.fill = purple_fill
            cell.font = header_font
            cell.alignment = center_align

        current_row += 1

        for record in week_groups[day]:

            for offset,value in enumerate(record):

                target_ws.cell(row=current_row,column=9+offset).value = value

            current_row += 1


# ------------------------------------------------------------
# SAVE FILE
# ------------------------------------------------------------

master_wb.save(OUTPUT_FILE)

files.download(OUTPUT_FILE)

print("Purple section updated successfully.")
