import pandas as pd
from datetime import datetime, timedelta, time
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO


def process(uploaded_file):

    df = pd.read_excel(uploaded_file, header=None)
    out = df.copy()

    # -----------------------------
    # COLUMN INDEX
    # -----------------------------
    MACL_AIRLINE = 0
    MACL_DAY = 1
    MACL_FLT = 3
    MACL_STA = 4
    MACL_STD = 5
    MACL_EFF = 6

    RAMIS_AIRLINE = 8
    RAMIS_DAY = 9
    RAMIS_FLT = 10
    RAMIS_STA = 11
    RAMIS_STD = 12
    RAMIS_EFF = 13

    RAMIS_START = 8

    NEW_FLT = 14
    NEW_STA = 15
    NEW_STD = 16
    NEW_EFF = 17
    STATUS  = 18
    COMMENT = 19

    # -----------------------------
    # TIME FUNCTIONS
    # -----------------------------
    def to_time(v):
        try:
            return datetime.strptime(str(v), "%H:%M").time()
        except:
            return None

    def base_flt(f):
        return str(f).split("-")[0]

    STA_START = time(4,45)
    STA_END   = time(15,45)

    STD_START = time(9,0)
    STD_END   = time(23,59)

    def sta_allowed(t):
        return t and STA_START <= t <= STA_END

    def std_allowed(t):
        return t and STD_START <= t <= STD_END

    # -----------------------------
    # BUILD RAMIS LOOKUP
    # -----------------------------
    ramis_rows = {}

    for i,r in df.iterrows():

        airline = str(r[RAMIS_AIRLINE]).strip().upper()
        day     = str(r[RAMIS_DAY]).strip().upper()
        flt     = str(r[RAMIS_FLT]).strip().upper()
        sta     = to_time(r[RAMIS_STA])
        std     = to_time(r[RAMIS_STD])

        if airline and day and flt and flt != "NAN":
            ramis_rows.setdefault(airline, []).append((day,flt,sta,std,i))

    used_ramis = set()

    # -----------------------------
    # MATCHING
    # -----------------------------
    for i,r in df.iterrows():

        airline  = str(r[MACL_AIRLINE]).strip().upper()
        macl_day = str(r[MACL_DAY]).strip().upper()
        macl_flt = str(r[MACL_FLT]).strip().upper()

        macl_sta = to_time(r[MACL_STA])
        macl_std = to_time(r[MACL_STD])

        macl_norm = macl_flt.replace("-","")
        macl_base = base_flt(macl_flt)

        found = False

        for r_day,r_flt,r_sta,r_std,r_i in ramis_rows.get(airline,[]):

            if r_i in used_ramis:
                continue

            if macl_day != r_day:
                continue

            ramis_norm = r_flt.replace("-","")
            ramis_base = base_flt(r_flt)

            if macl_norm == ramis_norm:
                out.iloc[i,RAMIS_START:] = df.iloc[r_i,RAMIS_START:]
                used_ramis.add(r_i)
                found = True
                break

            if macl_base == ramis_base:

                sta_ok = macl_sta == r_sta if sta_allowed(macl_sta) else True
                std_ok = macl_std == r_std if std_allowed(macl_std) else True

                if sta_ok and std_ok:
                    out.iloc[i,RAMIS_START:] = df.iloc[r_i,RAMIS_START:]
                    used_ramis.add(r_i)
                    found = True
                    break

        if not found:
            out.iloc[i,RAMIS_START:] = ""

    # -----------------------------
    # EXPORT TO MEMORY
    # -----------------------------
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        out.to_excel(writer, index=False, header=False)

    output.seek(0)

    return output
